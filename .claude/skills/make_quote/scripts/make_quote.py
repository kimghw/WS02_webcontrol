#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_quote.py — 장바구니·주문서·견적서 원본에서 뽑은 JSON 을 input.yaml 스키마로 검증하고 견적서로 렌더 (make_quote 스킬 실행체)

추출(원본 이미지·텍스트를 읽고 값을 뽑는 일)은 코드 에이전트가 input.yaml 의 키별 prompt 대로 한다. 이 스크립트는
그 결과 JSON 을 받아 스키마(타입·필수·enum·정합성 checks)로 검증하고 기본값을 채운 뒤 Markdown/CSV/JSON/XLSX
견적서로 쓴다. 검색 키워드: 견적서, 장바구니, input.yaml, validate, render.

서브커맨드 (출력은 항상 stdout JSON 1개, help 만 텍스트):
  schema   [--yaml F]                     input.yaml 의 키 트리(경로·type·required·default·label·prompt) + 원칙·checks·render
  template [--yaml F]                     채워 넣을 빈 JSON 골격(모든 키 null, items 1행)
  validate --in J [--yaml F] [--out-dir D]
                                          JSON 검증: 미지 키·필수 누락·타입·enum·checks. 기본값을 채운 data 도 반환
  render   --in J (--out F | --out-dir D) [--format md,json,csv,xlsx] [--force] [--yaml F]
                                          검증 후 견적서 렌더. --out 은 파일 1개(확장자로 형식 결정), --out-dir 은
                                          <견적번호>.<ext> 로 형식별 1개씩. 둘 다 없으면 content 로 반환(파일 안 씀).
                                          기존 파일은 --force 없이는 덮어쓰지 않음. 검증 error 가 있으면 렌더하지 않음
  help                                    이 도움말

`--in -` 는 stdin 에서 JSON 을 읽는다. JSON 의 `_` 로 시작하는 키는 메모용으로 무시한다.
종료코드: 0 정상 · 1 검증 error 또는 렌더 실패 · 2 인자·파일 오류
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from datetime import date
from pathlib import Path
from types import SimpleNamespace

try:  # Windows 콘솔(cp949) 에서도 한글 JSON 이 깨지지 않게
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

try:
    import yaml
except ImportError:  # pragma: no cover
    print(json.dumps({"ok": False, "error": "PyYAML 이 필요하다: pip install pyyaml"}, ensure_ascii=False))
    sys.exit(2)

SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_YAML = SKILL_DIR / "input.yaml"
DEFAULT_OUT_DIR = Path("quotes")
FORMATS = ("md", "json", "csv", "xlsx")
NUMERIC_TYPES = ("int", "money")

DATE_PATTERNS = [
    re.compile(r"^(\d{4})[-./년]\s*(\d{1,2})[-./월]\s*(\d{1,2})\s*일?$"),
    re.compile(r"^(\d{4})(\d{2})(\d{2})$"),
    re.compile(r"^(\d{2})[-./](\d{1,2})[-./](\d{1,2})$"),
]


# ----------------------------------------------------------------------------- 공통
def out(obj: dict, code: int = 0) -> None:
    print(json.dumps(obj, ensure_ascii=False, indent=2))
    sys.exit(code)


def load_schema(path: Path) -> dict:
    if not path.exists():
        out({"ok": False, "error": f"스키마 파일 없음: {path}"}, 2)
    try:
        with path.open("r", encoding="utf-8") as f:
            schema = yaml.safe_load(f) or {}
    except yaml.YAMLError as e:
        mark = getattr(e, "problem_mark", None)
        where = f" (line {mark.line + 1}, col {mark.column + 1})" if mark else ""
        out({"ok": False, "error": f"스키마 YAML 파싱 실패{where}: {getattr(e, 'problem', e)}. "
                                   "한 줄 prompt 에 ': ' 가 들어가면 'prompt: >-' 블록으로 쓴다", "yaml": str(path)}, 2)
    if "fields" not in schema:
        out({"ok": False, "error": f"스키마에 fields 가 없다: {path}"}, 2)
    return schema


def load_input(spec: str) -> dict:
    try:
        text = sys.stdin.read() if spec == "-" else Path(spec).read_text(encoding="utf-8-sig")
        data = json.loads(text)
    except FileNotFoundError:
        out({"ok": False, "error": f"입력 JSON 없음: {spec}"}, 2)
    except json.JSONDecodeError as e:
        out({"ok": False, "error": f"JSON 파싱 실패: {e}"}, 2)
    if not isinstance(data, dict):
        out({"ok": False, "error": "입력 JSON 최상위는 객체여야 한다"}, 2)
    return data


def enum_values(fdef: dict) -> dict:
    """values 가 리스트면 {값: 값}, 매핑이면 그대로."""
    vals = fdef.get("values") or {}
    if isinstance(vals, list):
        return {str(v): str(v) for v in vals}
    return {str(k): str(v) for k, v in vals.items()}


def parse_date(s: str) -> str | None:
    s = s.strip()
    for pat in DATE_PATTERNS:
        m = pat.match(s)
        if m:
            y, mo, d = m.groups()
            if len(y) == 2:
                y = "20" + y
            try:
                return date(int(y), int(mo), int(d)).isoformat()
            except ValueError:
                return None
    return None


def money_to_number(s: str):
    cleaned = re.sub(r"[^\d.\-]", "", s)
    if cleaned in ("", "-", "."):
        return None
    try:
        return int(cleaned) if re.fullmatch(r"-?\d+", cleaned) else float(cleaned)
    except ValueError:
        return None


# ----------------------------------------------------------------------------- 정규화·검증
class Report:
    def __init__(self) -> None:
        self.errors: list[dict] = []
        self.warnings: list[dict] = []
        self.normalized: list[dict] = []

    def err(self, path: str, code: str, msg: str) -> None:
        self.errors.append({"path": path, "code": code, "msg": msg})

    def warn(self, path: str, code: str, msg: str) -> None:
        self.warnings.append({"path": path, "code": code, "msg": msg})

    def norm(self, path: str, before, after) -> None:
        self.normalized.append({"path": path, "from": before, "to": after})


def coerce_scalar(value, fdef: dict, path: str, rep: Report):
    """스칼라 타입을 스키마에 맞게 강제 변환. 실패하면 error 를 남기고 원값 반환."""
    t = fdef.get("type", "string")
    if value is None:
        return None
    if t == "string":
        if not isinstance(value, str):
            rep.norm(path, value, str(value))
            return str(value)
        return value.strip()
    if t == "int":
        if isinstance(value, bool):
            rep.err(path, "type", "정수여야 한다")
            return value
        if isinstance(value, int):
            return value
        if isinstance(value, float) and value.is_integer():
            rep.norm(path, value, int(value))
            return int(value)
        if isinstance(value, str):
            n = money_to_number(value)
            if isinstance(n, int):
                rep.norm(path, value, n)
                return n
        rep.err(path, "type", f"정수여야 한다: {value!r}")
        return value
    if t == "money":
        if isinstance(value, bool):
            rep.err(path, "type", "금액은 숫자여야 한다")
            return value
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            return int(value) if value.is_integer() else value
        if isinstance(value, str):
            n = money_to_number(value)
            if n is not None:
                rep.norm(path, value, n)
                return n
        rep.err(path, "type", f"금액은 숫자여야 한다(원·쉼표 제거): {value!r}")
        return value
    if t == "bool":
        if isinstance(value, bool):
            return value
        if isinstance(value, str):
            low = value.strip().lower()
            if low in ("true", "yes", "y", "1", "예", "포함"):
                rep.norm(path, value, True)
                return True
            if low in ("false", "no", "n", "0", "아니오", "별도"):
                rep.norm(path, value, False)
                return False
        rep.err(path, "type", f"true/false 여야 한다: {value!r}")
        return value
    if t == "date":
        if isinstance(value, str):
            iso = parse_date(value)
            if iso is None:
                rep.err(path, "type", f"날짜(YYYY-MM-DD)여야 한다: {value!r}")
                return value
            if iso != value:
                rep.norm(path, value, iso)
            return iso
        rep.err(path, "type", f"날짜 문자열이어야 한다: {value!r}")
        return value
    if t == "url":
        if isinstance(value, str) and re.match(r"^https?://", value.strip()):
            return value.strip()
        rep.err(path, "type", f"http(s):// 로 시작하는 URL 이어야 한다: {value!r}")
        return value
    if t == "enum":
        allowed = enum_values(fdef)
        sval = str(value)
        if sval in allowed:
            return sval
        # 표시이름으로 준 경우(예 "장바구니") 값으로 되돌림
        for k, label in allowed.items():
            if label == sval:
                rep.norm(path, value, k)
                return k
        rep.err(path, "enum", f"허용값 {list(allowed)} 중 하나여야 한다: {value!r}")
        return value
    rep.err(path, "schema", f"알 수 없는 type: {t}")
    return value


def normalize_object(data, fields: dict, path: str, rep: Report) -> dict:
    """스키마 fields 순서대로 키를 정렬한 새 dict. 미지 키는 warning, `_` 키는 무시."""
    if data is None:
        data = {}
    if not isinstance(data, dict):
        rep.err(path or "(root)", "type", "객체여야 한다")
        return {}
    result: dict = {}
    for key, fdef in fields.items():
        p = f"{path}.{key}" if path else key
        val = data.get(key)
        t = fdef.get("type", "string")
        if t == "object":
            if val is None and not fdef.get("required"):
                result[key] = None
            else:  # required 객체가 비어 있으면 빈 객체로 만들어 하위 required 가 각각 보고되게 한다
                result[key] = normalize_object(val or {}, fdef.get("fields", {}), p, rep)
        elif t == "list":
            if val is None:
                result[key] = []
            elif not isinstance(val, list):
                rep.err(p, "type", "배열이어야 한다")
                result[key] = []
            else:
                item_fields = (fdef.get("item") or {}).get("fields", {})
                result[key] = [normalize_object(it, item_fields, f"{p}[{i}]", rep) for i, it in enumerate(val)]
        else:
            result[key] = coerce_scalar(val, fdef, p, rep)
    for key in data:
        if key not in fields and not str(key).startswith("_"):
            rep.warn(f"{path}.{key}" if path else key, "unknown_key", "스키마에 없는 키(무시됨)")
    return result


def next_seq(out_dir: Path, date_compact: str) -> str:
    pat = re.compile(rf"^Q-{date_compact}-(\d+)")
    mx = 0
    if out_dir.exists():
        for p in out_dir.iterdir():
            m = pat.match(p.name)
            if m:
                mx = max(mx, int(m.group(1)))
    return f"{mx + 1:02d}"


def apply_defaults(data: dict, fields: dict, path: str, ctx: dict, rep: Report) -> None:
    for key, fdef in fields.items():
        p = f"{path}.{key}" if path else key
        t = fdef.get("type", "string")
        if t == "object":
            if isinstance(data.get(key), dict):
                apply_defaults(data[key], fdef.get("fields", {}), p, ctx, rep)
        elif t == "list":
            item_fields = (fdef.get("item") or {}).get("fields", {})
            for i, it in enumerate(data.get(key) or []):
                apply_defaults(it, item_fields, f"{p}[{i}]", ctx, rep)
        elif data.get(key) is None and "default" in fdef:
            dv = fdef["default"]
            if isinstance(dv, str):
                if "{seq}" in dv and ctx.get("seq") is None:
                    ctx["seq"] = next_seq(ctx["out_dir"], ctx["date"])
                dv = dv.replace("{today}", ctx["today"]).replace("{date}", ctx["date"]).replace("{seq}", ctx.get("seq") or "")
            data[key] = dv
            rep.norm(p, None, dv)


def check_required(data: dict, fields: dict, path: str, rep: Report) -> None:
    for key, fdef in fields.items():
        p = f"{path}.{key}" if path else key
        t = fdef.get("type", "string")
        val = data.get(key)
        if t == "object":
            if isinstance(val, dict):
                check_required(val, fdef.get("fields", {}), p, rep)
            elif fdef.get("required"):
                rep.err(p, "required", "필수 객체가 없다")
        elif t == "list":
            items = val or []
            if fdef.get("required") and not items:
                rep.err(p, "required", "항목이 1개 이상 있어야 한다")
            elif fdef.get("min") and len(items) < int(fdef["min"]):
                rep.err(p, "min", f"항목이 {fdef['min']}개 이상이어야 한다")
            item_fields = (fdef.get("item") or {}).get("fields", {})
            for i, it in enumerate(items):
                check_required(it, item_fields, f"{p}[{i}]", rep)
        elif fdef.get("required") and val is None:
            rep.err(p, "required", "필수값이 비어 있다")


def _nz(x):
    return 0 if x is None else x


def _ns(d):
    return SimpleNamespace(**{k: v for k, v in (d or {}).items()})


def run_checks(data: dict, checks: list, rep: Report) -> list[dict]:
    results = []
    helpers = {"__builtins__": {}, "nz": _nz, "sum": sum, "abs": abs, "len": len, "min": min, "max": max, "round": round}

    def evaluate(rule: str, ns: dict):
        try:
            return bool(eval(rule, {**helpers, **ns})), None  # noqa: S307 — 스키마 파일은 로컬 신뢰 입력
        except Exception as e:  # None 연산 등
            return None, f"{type(e).__name__}: {e}"

    for c in checks or []:
        cid, rule, level, msg = c.get("id", "?"), c.get("rule", "True"), c.get("level", "warn"), c.get("message", "")
        targets = []
        if c.get("scope") == "item":
            for i, it in enumerate(data.get("items") or []):
                targets.append((f"items[{i}]", {**it}))
        else:
            targets.append(("", {"meta": _ns(data.get("meta")), "summary": _ns(data.get("summary")),
                                 "items": [_ns(i) for i in data.get("items") or []]}))
        for path, ns in targets:
            ok, err = evaluate(rule, ns)
            status = "pass" if ok else ("skipped" if ok is None else "fail")
            r = {"id": cid, "level": level, "status": status, "path": path, "message": msg}
            if err:
                r["reason"] = err
            results.append(r)
            if status == "fail":
                (rep.err if level == "error" else rep.warn)(path or "(quote)", f"check:{cid}", msg)
    return results


def validate_data(raw: dict, schema: dict, out_dir: Path) -> tuple[dict, Report, list]:
    rep = Report()
    fields = schema["fields"]
    data = normalize_object(raw, fields, "", rep)
    today = date.today().isoformat()
    qd = (data.get("meta") or {}).get("quote_date")
    ctx = {"today": today, "date": (qd if isinstance(qd, str) and parse_date(qd) else today).replace("-", ""), "out_dir": out_dir}
    apply_defaults(data, fields, "", ctx, rep)
    check_required(data, fields, "", rep)
    checks = run_checks(data, schema.get("checks") or [], rep)  # null 이 섞이면 해당 검사만 skipped
    return data, rep, checks


# ----------------------------------------------------------------------------- 렌더
def fmt_money(n, currency: str, schema: dict) -> str:
    if n is None:
        return ""
    s = f"{n:,}" if isinstance(n, int) else f"{n:,.2f}"
    tpl = ((schema.get("render") or {}).get("money_format") or {}).get(currency, "{n}")
    return tpl.replace("{n}", s)


def field_def(schema: dict, path: str) -> dict:
    cur = schema["fields"]
    parts = path.split(".")
    for i, part in enumerate(parts):
        fdef = cur.get(part, {})
        if i == len(parts) - 1:
            return fdef
        cur = fdef.get("fields") or (fdef.get("item") or {}).get("fields") or {}
    return {}


def display(value, fdef: dict) -> str:
    if value is None:
        return ""
    t = fdef.get("type", "string")
    if t == "enum":
        return enum_values(fdef).get(str(value), str(value))
    if t == "bool":
        return "예" if value else "아니오"
    return str(value)


def md_cell(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ")


def footer_line(meta: dict, schema: dict) -> str:
    tpl = (schema.get("render") or {}).get("footer") or ""
    segs = []
    for seg in tpl.split(" · "):
        keys = re.findall(r"\{(\w+)\}", seg)
        vals = {k: display(meta.get(k), field_def(schema, f"meta.{k}")) for k in keys}
        if keys and all(v == "" for v in vals.values()):
            continue
        for k, v in vals.items():
            seg = seg.replace("{" + k + "}", v)
        segs.append(re.sub(r"\s{2,}", " ", seg).strip())
    return " · ".join(s for s in segs if s)


def build_tables(data: dict, schema: dict) -> dict:
    """md/csv/xlsx 가 공유하는 표 데이터."""
    r = schema.get("render") or {}
    meta, items, summary = data.get("meta") or {}, data.get("items") or [], data.get("summary") or {}
    cur = meta.get("currency") or "KRW"

    header = []
    for k in r.get("header_keys") or []:
        fd = field_def(schema, f"meta.{k}")
        v = display(meta.get(k), fd)
        if v:
            header.append((fd.get("label", k), v))

    cols = [c for c in (r.get("item_columns") or []) if any(it.get(c) is not None for it in items)]
    col_defs = [(c, field_def(schema, f"items.{c}")) for c in cols]
    rows = []
    for i, it in enumerate(items, 1):
        row = [str(i)]
        for c, fd in col_defs:
            v = it.get(c)
            row.append(fmt_money(v, cur, schema) if fd.get("type") == "money" else display(v, fd))
        rows.append(row)
    item_header = ["No"] + [fd.get("label", c) for c, fd in col_defs]
    numeric = [False] + [fd.get("type") in NUMERIC_TYPES for _, fd in col_defs]

    srows = []
    for k in r.get("summary_rows") or []:
        v = summary.get(k)
        if v is None:
            continue
        fd = field_def(schema, f"summary.{k}")
        label = fd.get("label", k)
        if k == "discount":
            if summary.get("discount_note"):
                label += f" ({summary['discount_note']})"
            v_str = ("-" if v else "") + fmt_money(v, cur, schema)
        else:
            v_str = fmt_money(v, cur, schema)
        if k == "total":
            if meta.get("tax_included") is True:
                label += " (부가세 포함)"
            elif meta.get("tax_included") is False and summary.get("vat") is None:
                label += " (부가세 별도)"
        srows.append((label, v_str, k == "total"))

    return {"title": meta.get("title") or "견적서", "header": header, "item_header": item_header, "numeric": numeric,
            "rows": rows, "summary": srows, "footer": footer_line(meta, schema), "note": meta.get("note"),
            "items_label": schema["fields"].get("items", {}).get("label", "품목"),
            "summary_label": schema["fields"].get("summary", {}).get("label", "합계"),
            "raw_items": [(it, col_defs) for it in items], "currency": cur}


def render_md(t: dict) -> str:
    L = [f"# {t['title']}", ""]
    if t["header"]:
        L += ["| 항목 | 내용 |", "|---|---|"]
        L += [f"| {md_cell(k)} | {md_cell(v)} |" for k, v in t["header"]]
        L.append("")
    L += [f"## {t['items_label']}", ""]
    L.append("| " + " | ".join(t["item_header"]) + " |")
    L.append("|" + "|".join("---:" if n else "---" for n in t["numeric"]) + "|")
    for row in t["rows"]:
        L.append("| " + " | ".join(md_cell(c) for c in row) + " |")
    L += ["", f"## {t['summary_label']}", "", "| 구분 | 금액 |", "|---|---:|"]
    for label, v, is_total in t["summary"]:
        L.append(f"| **{md_cell(label)}** | **{v}** |" if is_total else f"| {md_cell(label)} | {v} |")
    L.append("")
    if t["footer"]:
        L.append(t["footer"])
    if t["note"]:
        L.append(f"비고: {t['note']}")
    return "\n".join(L).rstrip() + "\n"


def render_csv(t: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow([t["title"]])
    for k, v in t["header"]:
        w.writerow([k, v])
    w.writerow([])
    w.writerow(t["item_header"])
    for row in t["rows"]:
        w.writerow(row)
    w.writerow([])
    for label, v, _ in t["summary"]:
        w.writerow([label, v])
    if t["footer"]:
        w.writerow([])
        w.writerow([t["footer"]])
    if t["note"]:
        w.writerow([f"비고: {t['note']}"])
    return buf.getvalue()


def render_xlsx(t: dict, path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl 이 필요하다: pip install openpyxl")
    wb = Workbook()
    ws = wb.active
    ws.title = t["title"][:31] or "견적서"
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="EDEDED")
    ws.append([t["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in t["header"]:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])
    ws.append(t["item_header"])
    hdr_row = ws.max_row
    for c in range(1, len(t["item_header"]) + 1):
        cell = ws.cell(row=hdr_row, column=c)
        cell.font, cell.fill = bold, fill
    money_cols = {i + 1 for i, (c, fd) in enumerate(t["raw_items"][0][1], start=1) if fd.get("type") == "money"} if t["raw_items"] else set()
    int_cols = {i + 1 for i, (c, fd) in enumerate(t["raw_items"][0][1], start=1) if fd.get("type") == "int"} if t["raw_items"] else set()
    for n, (it, col_defs) in enumerate(t["raw_items"], 1):
        vals = [n] + [it.get(c) for c, _ in col_defs]
        ws.append(vals)
        r = ws.max_row
        for c in money_cols:
            ws.cell(row=r, column=c).number_format = "#,##0"
        for c in int_cols | money_cols | {1}:
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="right")
    ws.append([])
    for label, v_str, is_total in t["summary"]:
        n = money_to_number(v_str)
        ws.append([label, n if n is not None else v_str])
        r = ws.max_row
        ws.cell(row=r, column=2).number_format = "#,##0"
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="right")
        if is_total:
            ws.cell(row=r, column=1).font = ws.cell(row=r, column=2).font = bold
    if t["footer"]:
        ws.append([])
        ws.append([t["footer"]])
    if t["note"]:
        ws.append([f"비고: {t['note']}"])
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.row > 1:
                widths[cell.column] = max(widths.get(cell.column, 0), min(60, len(str(cell.value)) * 1.1 + 2))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(8, w)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ----------------------------------------------------------------------------- 서브커맨드
def cmd_schema(a) -> None:
    schema = load_schema(Path(a.yaml))
    flat = []

    def walk(fields: dict, path: str):
        for key, fdef in fields.items():
            p = f"{path}.{key}" if path else key
            entry = {"path": p, "type": fdef.get("type", "string"), "required": bool(fdef.get("required")),
                     "label": fdef.get("label"), "prompt": (fdef.get("prompt") or "").strip()}
            for opt in ("default", "values", "example", "min"):
                if opt in fdef:
                    entry[opt] = fdef[opt]
            flat.append(entry)
            if fdef.get("type") == "object":
                walk(fdef.get("fields", {}), p)
            elif fdef.get("type") == "list":
                walk((fdef.get("item") or {}).get("fields", {}), p + "[]")

    walk(schema["fields"], "")
    out({"ok": True, "action": "schema", "yaml": str(Path(a.yaml).resolve()), "version": schema.get("version"),
         "extraction_prompt": ((schema.get("extraction") or {}).get("prompt") or "").strip(),
         "fields": flat, "checks": schema.get("checks") or [], "render": schema.get("render") or {}})


def skeleton(fields: dict) -> dict:
    d = {}
    for key, fdef in fields.items():
        t = fdef.get("type", "string")
        if t == "object":
            d[key] = skeleton(fdef.get("fields", {}))
        elif t == "list":
            d[key] = [skeleton((fdef.get("item") or {}).get("fields", {}))]
        else:
            d[key] = None
    return d


def cmd_template(a) -> None:
    schema = load_schema(Path(a.yaml))
    out({"ok": True, "action": "template", "yaml": str(Path(a.yaml).resolve()), "template": skeleton(schema["fields"])})


def report_dict(rep: Report, checks: list) -> dict:
    return {"errors": rep.errors, "warnings": rep.warnings, "normalized": rep.normalized, "checks": checks}


def cmd_validate(a) -> None:
    schema = load_schema(Path(a.yaml))
    raw = load_input(a.inp)
    data, rep, checks = validate_data(raw, schema, Path(a.out_dir) if a.out_dir else DEFAULT_OUT_DIR)
    out({"ok": not rep.errors, "action": "validate", "quote_no": (data.get("meta") or {}).get("quote_no"),
         "item_count": len(data.get("items") or []), "total": (data.get("summary") or {}).get("total"),
         **report_dict(rep, checks), "data": data}, 0 if not rep.errors else 1)


def cmd_render(a) -> None:
    schema = load_schema(Path(a.yaml))
    raw = load_input(a.inp)
    out_path = Path(a.out) if a.out else None
    out_dir = Path(a.out_dir) if a.out_dir else (out_path.parent if out_path else DEFAULT_OUT_DIR)
    if out_path and a.out_dir:
        out({"ok": False, "error": "--out 과 --out-dir 는 함께 쓸 수 없다"}, 2)

    if out_path:
        ext = out_path.suffix.lstrip(".").lower()
        if a.format and a.format != ext:
            out({"ok": False, "error": f"--out 확장자({ext})와 --format({a.format})이 다르다"}, 2)
        formats = [ext or "md"]
    else:
        formats = [f.strip().lower() for f in (a.format or "md").split(",") if f.strip()]
    bad = [f for f in formats if f not in FORMATS]
    if bad:
        out({"ok": False, "error": f"지원하지 않는 형식 {bad}. 가능: {list(FORMATS)}"}, 2)
    if not out_path and not a.out_dir and "xlsx" in formats:
        out({"ok": False, "error": "xlsx 는 --out 또는 --out-dir 가 필요하다(내용을 inline 으로 돌려줄 수 없음)"}, 2)

    data, rep, checks = validate_data(raw, schema, out_dir)
    base = {"action": "render", "quote_no": (data.get("meta") or {}).get("quote_no"),
            "item_count": len(data.get("items") or []), "total": (data.get("summary") or {}).get("total"),
            **report_dict(rep, checks)}
    if rep.errors:
        out({"ok": False, "reason": "검증 error 가 있어 렌더하지 않았다", **base}, 1)

    t = build_tables(data, schema)
    content = {"md": lambda: render_md(t), "csv": lambda: render_csv(t),
               "json": lambda: json.dumps(data, ensure_ascii=False, indent=2) + "\n"}

    if not out_path and not a.out_dir:  # dry-run: 내용만 돌려준다
        out({"ok": True, "written": [], "content": {f: content[f]() for f in formats}, **base})

    targets = [(f, out_path if out_path else out_dir / f"{base['quote_no'] or 'quote'}.{f}") for f in formats]
    clashes = [str(p) for _, p in targets if p.exists()]
    if clashes and not a.force:
        out({"ok": False, "reason": "이미 있는 파일. 덮어쓰려면 --force", "exists": clashes, **base}, 1)
    written = []
    for f, p in targets:
        p.parent.mkdir(parents=True, exist_ok=True)
        try:
            if f == "xlsx":
                render_xlsx(t, p)
            else:
                p.write_text(content[f](), encoding="utf-8-sig" if f == "csv" else "utf-8", newline="\n")
        except RuntimeError as e:
            out({"ok": False, "error": str(e), "written": written, **base}, 1)
        written.append({"format": f, "path": str(p.resolve()), "overwritten": str(p) in clashes})
    out({"ok": True, "written": written, **base})


# ----------------------------------------------------------------------------- CLI
def main(argv: list[str]) -> None:
    if not argv or argv[0] in ("help", "-h", "--help"):
        print(__doc__.strip())
        sys.exit(0)
    ap = argparse.ArgumentParser(prog="make_quote.py", add_help=False)
    sub = ap.add_subparsers(dest="cmd")
    for name in ("schema", "template", "validate", "render"):
        sp = sub.add_parser(name, add_help=False)
        sp.add_argument("--yaml", default=str(DEFAULT_YAML))
        if name in ("validate", "render"):
            sp.add_argument("--in", dest="inp", required=True)
            sp.add_argument("--out-dir", dest="out_dir")
        if name == "render":
            sp.add_argument("--out")
            sp.add_argument("--format")
            sp.add_argument("--force", action="store_true")
    try:
        a = ap.parse_args(argv)
    except SystemExit:
        out({"ok": False, "error": "인자 오류. `help` 참고"}, 2)
    {"schema": cmd_schema, "template": cmd_template, "validate": cmd_validate, "render": cmd_render}[a.cmd](a)


if __name__ == "__main__":
    main(sys.argv[1:])
